#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
limpiatexto_cli v1.6 — Extrae texto "crudo" de documentos y lo guarda SIN formato en .txt

Formatos soportados (si las dependencias están instaladas):
  - Texto: .txt
  - Office/ODF: .docx, .odt, .ods, .xlsx, .csv, .pptx
  - Web/ebooks: .html/.htm, .epub
  - Rich Text: .rtf
  - PDF: .pdf  (texto; con OCR opcional si es escaneado)
  - DOC antiguo: .doc (vía 'antiword' si está disponible)
  - Imágenes a texto (OCR): .png/.jpg/.jpeg/.tif/.tiff (si hay pytesseract)

Notas OCR:
  - Para PDF escaneado o imágenes, si defines RAWTEXT_OCR=1 se intentará OCR automático.
  - Requiere: pip install pytesseract pdf2image pillow  +  paquete de sistema 'tesseract-ocr'
  - Para PDF con pdf2image también es recomendable 'poppler-utils'.

Dependencias Python opcionales (instala sólo lo que necesites):
  pip install python-docx odfpy pdfminer.six striprtf beautifulsoup4 ebooklib openpyxl python-pptx
  # OCR opcional:
  pip install pytesseract pdf2image pillow

Paquetes de sistema útiles (Debian/Ubuntu):
  sudo apt-get install antiword tesseract-ocr poppler-utils

- Sin argumentos, pregunta la ruta del archivo a convertir.
- Con rutas/"globs", los procesa en lote.
- La salida se limpia (espacios/tabs/nbsp, líneas vacías) y se guarda con nombre "sin formato-<original>.txt".
- Guardado robusto y "Windows-friendly" con fallback (Documentos/Temp) y sufijos -1, -2, ... si existe/bloqueado.
"""

from __future__ import annotations
import sys
import os
import re
import glob
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Optional, List

# --- Información del programa ---
__program_name__ = "limpiatexto"
__version__ = "1.6"
__date__ = "2025-10-22"
__author__ = "Oscar Gimenez Blasco"

# ==============================
#   Imports opcionales
# ==============================
# DOCX
try:
    import docx
    from docx.text.paragraph import Paragraph as DocxParagraph
    from docx.table import Table as DocxTable
    from docx.oxml.text.paragraph import CT_P as DOCX_CT_P
    from docx.oxml.table import CT_Tbl as DOCX_CT_TBL
except Exception:
    print("WARN: Falta 'python-docx' para .docx (pip install python-docx)", file=sys.stderr)
    docx = None  # type: ignore
    DOCX_CT_P = DOCX_CT_TBL = DocxParagraph = DocxTable = None  # type: ignore

# ODF (odt/ods)
try:
    from odf.opendocument import load as odf_load
    from odf.text import P as ODF_P
    from odf.teletype import extractText as odf_extract_text
    from odf.table import (
        Table as ODF_Table,
        TableRow as ODF_TableRow,
        TableCell as ODF_TableCell,
        CoveredTableCell as ODF_CoveredTableCell,
    )
except Exception:
    print("WARN: Falta 'odfpy' para .odt/.ods (pip install odfpy)", file=sys.stderr)
    odf_load = ODF_P = odf_extract_text = None  # type: ignore
    ODF_Table = ODF_TableRow = ODF_TableCell = ODF_CoveredTableCell = None  # type: ignore

# PDF texto
try:
    from pdfminer.high_level import extract_text as pdf_extract_text
except Exception:
    print("WARN: Falta 'pdfminer.six' para .pdf (pip install pdfminer.six)", file=sys.stderr)
    pdf_extract_text = None  # type: ignore

# OCR opcional (PDF escaneado / imágenes)
try:
    import pytesseract  # type: ignore
except Exception:
    pytesseract = None  # type: ignore

try:
    from pdf2image import convert_from_path  # type: ignore
except Exception:
    convert_from_path = None  # type: ignore

try:
    from PIL import Image  # type: ignore
except Exception:
    Image = None  # type: ignore

# HTML/EPUB
try:
    from bs4 import BeautifulSoup  # type: ignore
except Exception:
    BeautifulSoup = None  # type: ignore

try:
    from ebooklib import epub  # type: ignore
except Exception:
    epub = None  # type: ignore

# XLSX/CSV
try:
    from openpyxl import load_workbook  # type: ignore
except Exception:
    load_workbook = None  # type: ignore

import csv

# PPTX
try:
    from pptx import Presentation  # type: ignore
except Exception:
    Presentation = None  # type: ignore

# RTF
try:
    from striprtf.striprtf import rtf_to_text  # type: ignore
except Exception:
    rtf_to_text = None  # type: ignore


# ==============================
#   Utilidades de limpieza
# ==============================

def limpiar_texto_crudo(texto_sucio: str) -> str:
    """Normaliza el texto: tabs/nbsp -> espacios, colapsa espacios, quita vacías."""
    texto = texto_sucio.replace("\t", " ").replace("\xa0", " ")
    lineas = texto.splitlines()
    lineas_limpias: List[str] = []
    for linea in lineas:
        s = re.sub(r" +", " ", linea.strip())
        if s:
            lineas_limpias.append(s)
    return "\n".join(lineas_limpias)


# ==============================
#   Lectura DOCX (con tablas)
# ==============================

def _docx_iter_block_items(document_or_cell):
    """Itera párrafos y tablas manteniendo orden (documento o celdas)."""
    if document_or_cell is None:
        return
    if hasattr(document_or_cell, "element") and hasattr(document_or_cell.element, "body"):
        parent_elm = document_or_cell.element.body
        parent = document_or_cell
    elif hasattr(document_or_cell, "_tc"):
        parent_elm = document_or_cell._tc
        parent = document_or_cell
    else:
        return
    for child in parent_elm.iterchildren():
        if isinstance(child, DOCX_CT_P):
            yield DocxParagraph(child, parent)
        elif isinstance(child, DOCX_CT_TBL):
            yield DocxTable(child, parent)

def _docx_table_to_lines(tbl: 'DocxTable') -> List[str]:  # type: ignore[name-defined]
    lines: List[str] = []
    for row in tbl.rows:
        cells_text: List[str] = []
        for cell in row.cells:
            cell_lines: List[str] = []
            for block in _docx_iter_block_items(cell):
                if isinstance(block, DocxParagraph):
                    txt = (block.text or "").strip()
                    if txt:
                        cell_lines.append(txt)
                elif isinstance(block, DocxTable):
                    nested = _docx_table_to_lines(block)
                    if nested:
                        cell_lines.append(" ".join(nested))
            joined = " ".join(cell_lines).strip()
            if joined:
                cells_text.append(joined)
        if cells_text:
            lines.append(" ".join(cells_text))
    return lines

def leer_docx(ruta_archivo: str) -> str:
    if docx is None:
        raise RuntimeError("python-docx no está instalado.")
    d = docx.Document(ruta_archivo)
    out_lines: List[str] = []
    for block in _docx_iter_block_items(d):
        if isinstance(block, DocxParagraph):
            txt = (block.text or "").strip()
            if txt:
                out_lines.append(txt)
        elif isinstance(block, DocxTable):
            out_lines.extend(_docx_table_to_lines(block))
    return "\n".join(out_lines)


# ==============================
#   Lectura ODT/ODS
# ==============================

def leer_odt(ruta_archivo: str) -> str:
    if odf_load is None:
        raise RuntimeError("odfpy no está instalado.")
    doc = odf_load(ruta_archivo)
    parrafos = doc.getElementsByType(ODF_P)
    lines: List[str] = []
    for p in parrafos:
        txt = (odf_extract_text(p) or "").strip()
        if txt:
            lines.append(txt)
    return "\n".join(lines)

def leer_ods(ruta_archivo: str) -> str:
    """Extrae texto por filas: cada fila -> línea; celdas unidas por espacio."""
    if odf_load is None:
        raise RuntimeError("odfpy no está instalado.")
    doc = odf_load(ruta_archivo)
    lines: List[str] = []
    for table in doc.getElementsByType(ODF_Table):
        for row in table.getElementsByType(ODF_TableRow):
            cell_texts: List[str] = []
            for cell in row.childNodes:
                if not isinstance(cell, (ODF_TableCell, ODF_CoveredTableCell)):
                    continue
                ps = getattr(cell, "getElementsByType", lambda *_: [])(ODF_P)
                pieces = []
                for p in ps:
                    t = (odf_extract_text(p) or "").strip()
                    if t:
                        pieces.append(t)
                if pieces:
                    cell_texts.append(" ".join(pieces))
            if cell_texts:
                lines.append(" ".join(cell_texts))
    return "\n".join(lines)


# ==============================
#   Lectura XLSX / CSV
# ==============================

def leer_xlsx(ruta_archivo: str) -> str:
    if load_workbook is None:
        raise RuntimeError("openpyxl no está instalado.")
    wb = load_workbook(ruta_archivo, read_only=True, data_only=True)
    lines: List[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
            if cells:
                lines.append(" ".join(cells))
    return "\n".join(lines)

def leer_csv(ruta_archivo: str) -> str:
    def _read(enc: str) -> str:
        with open(ruta_archivo, "r", encoding=enc, newline="") as f:
            reader = csv.reader(f)
            return "\n".join(" ".join(cell.strip() for cell in row if cell.strip()) for row in reader)
    try:
        return _read("utf-8")
    except UnicodeDecodeError:
        return _read("latin-1")


# ==============================
#   Lectura HTML / EPUB / RTF / PPTX
# ==============================

def _html_to_text(html: str) -> str:
    if BeautifulSoup is None:
        raise RuntimeError("beautifulsoup4 no está instalado.")
    soup = BeautifulSoup(html, "html.parser")
    # Elimina elementos irrelevantes
    for tag in soup(["script", "style", "noscript", "header", "footer", "nav", "svg"]):
        tag.decompose()
    text = soup.get_text("\n")
    return text

def leer_html(ruta_archivo: str) -> str:
    try:
        with open(ruta_archivo, "r", encoding="utf-8") as f:
            html = f.read()
    except UnicodeDecodeError:
        with open(ruta_archivo, "r", encoding="latin-1") as f:
            html = f.read()
    return _html_to_text(html)

def leer_epub(ruta_archivo: str) -> str:
    if epub is None or BeautifulSoup is None:
        raise RuntimeError("ebooklib/bs4 no están instalados.")
    book = epub.read_epub(ruta_archivo)
    parts: List[str] = []
    for item in book.get_items():
        if item.get_type() == epub.ITEM_DOCUMENT:  # type: ignore[attr-defined]
            html = item.get_content().decode(errors="ignore")
            txt = _html_to_text(html)
            if txt:
                parts.append(txt)
    return "\n".join(parts)

def leer_rtf(ruta_archivo: str) -> str:
    if rtf_to_text is None:
        raise RuntimeError("striprtf no está instalado.")
    with open(ruta_archivo, "r", encoding="utf-8", errors="ignore") as f:
        data = f.read()
    return rtf_to_text(data)

def leer_pptx(ruta_archivo: str) -> str:
    if Presentation is None:
        raise RuntimeError("python-pptx no está instalado.")
    prs = Presentation(ruta_archivo)
    lines: List[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = (p.text or "").strip()
                    if t:
                        lines.append(t)
    return "\n".join(lines)


# ==============================
#   Lectura PDF
# ==============================

def _pdf_text(ruta_archivo: str) -> str:
    if pdf_extract_text is None:
        raise RuntimeError("pdfminer.six no está instalado.")
    # Nota: pdfminer puede tardar en PDFs grandes; es el más fiable para texto.
    try:
        txt = pdf_extract_text(ruta_archivo) or ""
    except Exception as e:
        raise RuntimeError(f"pdfminer error: {e!r}")
    return txt

def _have_ocr_stack() -> bool:
    return (pytesseract is not None) and (Image is not None)

def _pdf_ocr_text(ruta_archivo: str) -> str:
    if not _have_ocr_stack():
        raise RuntimeError("OCR no disponible (instala pytesseract y pillow).")
    images: List['Image.Image'] = []  # type: ignore[name-defined]
    if convert_from_path is not None:
        try:
            # Requiere poppler-utils para mejor compatibilidad
            images = convert_from_path(ruta_archivo)
        except Exception as e:
            print(f"WARN: pdf2image falló ({e!r}); intentando con 'pdftoppm' externo si existe...", file=sys.stderr)
    if not images:
        # Fallback a 'pdftoppm' si existe
        pdftoppm = shutil.which('pdftoppm')
        if not pdftoppm:
            raise RuntimeError("No se pudo rasterizar PDF (instala poppler-utils o pdf2image).")
        with tempfile.TemporaryDirectory() as tmpd:
            prefix = Path(tmpd) / "page"
            cmd = [pdftoppm, "-r", "200", "-png", ruta_archivo, str(prefix)]
            subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            for png in sorted(Path(tmpd).glob("page-*.png")):
                if Image is None:
                    break
                images.append(Image.open(png))
    out: List[str] = []
    for im in images:
        try:
            txt = pytesseract.image_to_string(im) if pytesseract else ""
            if txt:
                out.append(txt)
        finally:
            try:
                im.close()
            except Exception:
                pass
    return "\n".join(out)

def leer_pdf(ruta_archivo: str) -> str:
    text = _pdf_text(ruta_archivo)
    # Si no hay texto (posible PDF escaneado), intenta OCR si activado
    if len(text.strip()) < 20 and os.getenv("RAWTEXT_OCR", "0") in ("1", "true", "True"):
        try:
            ocr_txt = _pdf_ocr_text(ruta_archivo)
            if ocr_txt.strip():
                return ocr_txt
        except Exception as e:
            print(f"WARN: OCR PDF falló: {e}", file=sys.stderr)
    return text


# ==============================
#   Lectura DOC (antiguo) con antiword
# ==============================

def leer_doc_via_antiword(ruta_archivo: str) -> Optional[str]:
    anti = shutil.which('antiword')
    if not anti:
        return None
    mapfile = "/usr/share/antiword/UTF-8.txt"
    cmd = [anti]
    if os.path.exists(mapfile):
        cmd += ["-m", mapfile]
    cmd.append(ruta_archivo)
    try:
        res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=60)
        if res.returncode == 0 and res.stdout.strip():
            return res.stdout
        else:
            print(f"WARN: antiword no pudo extraer texto (rc={res.returncode}, err={res.stderr.strip()[:120]})", file=sys.stderr)
            return None
    except Exception as e:
        print(f"WARN: antiword lanzó excepción: {e!r}", file=sys.stderr)
        return None


# ==============================
#   Imágenes -> OCR directo
# ==============================

def leer_imagen_ocr(ruta_archivo: str) -> str:
    if not _have_ocr_stack():
        raise RuntimeError("OCR no disponible (pytesseract/pillow)")
    if Image is None:
        raise RuntimeError("Pillow no está instalado")
    im = Image.open(ruta_archivo)
    try:
        return pytesseract.image_to_string(im) if pytesseract else ""
    finally:
        try:
            im.close()
        except Exception:
            pass


# ==============================
#   Multiplexor de lectura
# ==============================

def _read_txt(ruta: str) -> str:
    try:
        with open(ruta, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(ruta, "r", encoding="latin-1") as f:
            return f.read()


def leer_archivo(ruta_archivo: str) -> Optional[str]:
    _, extension = os.path.splitext(ruta_archivo)
    ext = extension.lower()
    try:
        if ext == ".txt":
            return _read_txt(ruta_archivo)
        elif ext == ".docx":
            return leer_docx(ruta_archivo)
        elif ext == ".odt":
            return leer_odt(ruta_archivo)
        elif ext == ".ods":
            return leer_ods(ruta_archivo)
        elif ext == ".xlsx":
            return leer_xlsx(ruta_archivo)
        elif ext == ".csv":
            return leer_csv(ruta_archivo)
        elif ext in (".html", ".htm"):
            return leer_html(ruta_archivo)
        elif ext == ".epub":
            return leer_epub(ruta_archivo)
        elif ext == ".rtf":
            return leer_rtf(ruta_archivo)
        elif ext == ".pptx":
            return leer_pptx(ruta_archivo)
        elif ext == ".pdf":
            return leer_pdf(ruta_archivo)
        elif ext == ".doc":  # antiguo
            via = leer_doc_via_antiword(ruta_archivo)
            return via if via is not None else "formato_doc_no_soportado"
        elif ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            if os.getenv("RAWTEXT_OCR", "0") in ("1", "true", "True"):
                return leer_imagen_ocr(ruta_archivo)
            else:
                return "ocr_desactivado_para_imagen"
        else:
            return None
    except Exception as e:
        return f"error_lectura:{e!r}"


# ==============================
#   Guardado robusto (Windows-friendly)
# ==============================

_ILLEGAL_WIN_CHARS = r'<>:"/\\|?*'


def sanitize_filename(name: str) -> str:
    trans = {ord(c): "_" for c in _ILLEGAL_WIN_CHARS}
    name = name.translate(trans)
    name = re.sub(r"\s+", " ", name).strip()
    name = re.sub(r"[ \.]+$", "", name)
    reserved = re.compile(r"^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(\.|$)", re.IGNORECASE)
    if reserved.match(name):
        name = "_" + name
    if len(name) > 240:
        name = re.sub(r"[ \.]+$", "", name[:240])
    return name


def _dir_is_writable(path_dir: Path) -> bool:
    try:
        path_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=str(path_dir), delete=True) as _:
            pass
        return True
    except Exception:
        return False


def _default_fallback_dir() -> Path:
    home = Path.home()
    for d in [home / "Documents", home / "Documentos", Path(tempfile.gettempdir())]:
        if _dir_is_writable(d):
            return d
    return Path.cwd()


def build_output_path(input_path: Path, base_output_name: str) -> Path:
    """Intenta guardar junto al archivo original; si no, usa fallback. Añade sufijos -1, -2..."""
    safe_base = sanitize_filename(base_output_name)
    candidate_dirs = [input_path.parent]
    if not _dir_is_writable(input_path.parent):
        candidate_dirs = [_default_fallback_dir()]

    for out_dir in candidate_dirs:
        n = 0
        while True:
            suffix = f"-{n}" if n else ""
            out_path = out_dir / f"{safe_base}{suffix}.txt"
            try:
                with open(out_path, "x", encoding="utf-8") as f:
                    f.write("")
                return out_path
            except FileExistsError:
                n += 1
                continue
            except PermissionError:
                if out_dir == input_path.parent:
                    alt = _default_fallback_dir()
                    if alt != out_dir:
                        candidate_dirs = [alt]
                        break
                n += 1
                continue
            except OSError:
                n += 1
                continue
    # Último recurso
    return Path(tempfile.gettempdir()) / f"{safe_base}.txt"


# ==============================
#   Procesamiento CLI
# ==============================

def procesar_archivo_cli(path_str: str) -> bool:
    in_path = Path(path_str).expanduser()
    if not in_path.exists():
        print(f"❌ ERROR: No existe: {path_str}", file=sys.stderr)
        return False

    print(f"📄 Procesando: {in_path.name}")
    contenido_original = leer_archivo(str(in_path))

    if contenido_original == "formato_doc_no_soportado":
        print("   ⚠️  Formato .doc (antiguo) no soportado. Instala 'antiword' para soporte.", file=sys.stderr)
        return False
    if contenido_original == "ocr_desactivado_para_imagen":
        print("   ⚠️  Imagen detectada. Activa OCR con RAWTEXT_OCR=1 si quieres extraer texto.", file=sys.stderr)
        return False
    if contenido_original is None:
        print("   ❌ Formato no reconocido.", file=sys.stderr)
        return False
    if isinstance(contenido_original, str) and contenido_original.startswith("error_lectura:"):
        print(f"   ❌ Error al leer: {contenido_original}", file=sys.stderr)
        return False

    contenido_limpio = limpiar_texto_crudo(contenido_original)
    base_nombre = f"sin formato-{in_path.stem}"
    out_path = build_output_path(in_path, base_nombre)

    try:
        with open(out_path, "w", encoding="utf-8") as f_out:
            f_out.write(contenido_limpio)
        print(f"   ✅ Guardado: {out_path}")
        return True
    except PermissionError:
        # Intento final en fallback
        fallback_dir = _default_fallback_dir()
        dummy_in = fallback_dir / "dummy.txt"
        out_path2 = build_output_path(dummy_in, base_nombre)
        try:
            with open(out_path2, "w", encoding="utf-8") as f_out:
                f_out.write(contenido_limpio)
            print("   ⚠️  Sin permisos en la carpeta original. Guardado en fallback:")
            print(f"      {out_path2}")
            return True
        except Exception as e2:
            print(f"   ❌ Error de permisos incluso en fallback: {e2!r}", file=sys.stderr)
            return False
    except Exception as e:
        print(f"   ❌ Error inesperado al guardar: {e!r}", file=sys.stderr)
        return False


def _expand_globs(paths: List[str]) -> List[str]:
    """En Windows CMD los * no se expanden: lo hacemos aquí."""
    out: List[str] = []
    for p in paths:
        matches = glob.glob(p)
        out.extend(matches if matches else [p])
    return out


def main(argv: List[str] | None = None) -> int:
    if argv is None:
        argv = sys.argv
    args = argv[1:]

    # Ayuda mínima sin argparse
    if args and args[0] in ("-h", "--help"):
        print(
            "Uso: sabbat-rawtext [ficheros...]\n"
            " - Sin argumentos: pregunta la ruta por consola.\n"
            " - Soporta: .txt .docx .odt .ods .xlsx .csv .html/.htm .epub .rtf .pptx .pdf .doc (antiword)\n"
            " - Imágenes (.png/.jpg/.jpeg/.tif/.tiff) vía OCR si RAWTEXT_OCR=1.\n"
            " - Globs (comodines) aceptados; en Windows se expanden internamente.\n"
        )
        return 0

    if not args:
        try:
            path = input("Ruta del archivo a convertir (ver --help para formatos): ").strip('"').strip()
        except EOFError:
            return 1
        if not path:
            print("Nada que hacer.", file=sys.stderr)
            return 1
        targets = [path]
    else:
        targets = _expand_globs(args)

    ok = True
    for target in targets:
        ok &= procesar_archivo_cli(target)

    print("\n🎉 Proceso terminado.")
    return 0 if ok else 1


def cli_main() -> int:
    """Wrapper de compatibilidad para entry points antiguos."""
    return main()


if __name__ == "__main__":
    raise SystemExit(main())

